"""
AgriSense: Deep Learning Honors Project Report Generator
Constructs a comprehensive, rigorous academic project report strictly adhering to the
Shivaji University / DKTE template and the DL Honors Project Report Contents guideline.
Outputs: AgriSense_Honors_Project_Report.docx
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "AgriSense_Honors_Project_Report.docx")
FIG_DIR = os.path.join(os.path.dirname(__file__), "report_figures")

def set_cell_background(cell, fill_hex):
    """Sets background shading of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    """Sets light subtle borders on a table for clean academic look."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="2E4053"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_callout(doc, text, title="NOTE"):
    """Adds a stylish callout block for notes or mathematical takeaways."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.2)
    set_cell_background(cell, "F2F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    tcPr = cell._tc.get_or_add_tcPr()
    border = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="1F4E79"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(border)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"[{title}] ")
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(10)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    run_b = p.add_run(text)
    run_b.font.name = "Times New Roman"
    run_b.font.size = Pt(10)
    run_b.font.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_paragraph(p, space_before=0, space_after=6, line_spacing=1.3, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align

def add_body_p(doc, text, bold_prefix=None, space_after=6):
    p = doc.add_paragraph()
    format_paragraph(p, space_after=space_after)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    return p

def add_heading_1(doc, text):
    """Chapter heading."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=16, space_after=12, align=WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)
    return p

def add_heading_2(doc, text):
    """Section heading."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=14, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(13.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    return p

def add_heading_3(doc, text):
    """Subsection heading."""
    p = doc.add_paragraph()
    format_paragraph(p, space_before=10, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x34, 0x49, 0x5E)
    return p

def add_equation(doc, eq_text, eq_number):
    """Adds a centered mathematical equation with right-aligned numbering."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell_eq = table.rows[0].cells[0]
    cell_num = table.rows[0].cells[1]
    cell_eq.width = Inches(5.4)
    cell_num.width = Inches(0.8)
    set_cell_margins(cell_eq, top=60, bottom=60, left=60, right=60)
    set_cell_margins(cell_num, top=60, bottom=60, left=60, right=60)

    p_eq = cell_eq.paragraphs[0]
    format_paragraph(p_eq, space_before=2, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_eq = p_eq.add_run(eq_text)
    r_eq.font.name = "Times New Roman"
    r_eq.font.size = Pt(11)
    r_eq.font.italic = True

    p_num = cell_num.paragraphs[0]
    format_paragraph(p_num, space_before=2, space_after=2, align=WD_ALIGN_PARAGRAPH.RIGHT)
    r_num = p_num.add_run(f"({eq_number})")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(11)
    r_num.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_figure(doc, img_filename, fig_num_title, width_inches=6.2):
    """Inserts a figure image with standard academic caption below."""
    img_path = os.path.join(FIG_DIR, img_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        format_paragraph(p_img, space_before=10, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
        p_img.add_run().add_picture(img_path, width=Inches(width_inches))
        
        p_cap = doc.add_paragraph()
        format_paragraph(p_cap, space_before=3, space_after=14, align=WD_ALIGN_PARAGRAPH.CENTER)
        r_num = p_cap.add_run(fig_num_title.split(":")[0] + ":")
        r_num.font.name = "Times New Roman"
        r_num.font.size = Pt(10.5)
        r_num.font.bold = True
        r_txt = p_cap.add_run(fig_num_title[len(fig_num_title.split(":")[0]) + 1:])
        r_txt.font.name = "Times New Roman"
        r_txt.font.size = Pt(10.5)
        r_txt.font.italic = True
    else:
        print(f"[!] Warning: Figure not found at {img_path}")

def add_table_with_caption(doc, caption, headers, rows_data, col_widths=None):
    """Adds an academic styled table with caption above."""
    p_cap = doc.add_paragraph()
    format_paragraph(p_cap, space_before=10, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    r_cap_lbl = p_cap.add_run(caption.split(":")[0] + ":")
    r_cap_lbl.font.name = "Times New Roman"
    r_cap_lbl.font.size = Pt(10.5)
    r_cap_lbl.font.bold = True
    r_cap_txt = p_cap.add_run(caption[len(caption.split(":")[0]) + 1:])
    r_cap_txt.font.name = "Times New Roman"
    r_cap_txt.font.size = Pt(10.5)

    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    # Format header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "2C3E50")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in hdr_cells[i].paragraphs[0].runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Format data rows
    for r_idx, row in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=140, right=140)
            p = row_cells[c_idx].paragraphs[0]
            format_paragraph(p, space_before=2, space_after=2, line_spacing=1.15)
            # Center numeric/short fields, left align long descriptive fields
            if len(str(val)) < 12 and not any(c.islower() and len(str(val)) > 8 for c in str(val)):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)

    if col_widths:
        for row in table.rows:
            for c_idx, w in enumerate(col_widths):
                row.cells[c_idx].width = Inches(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

print("[*] Builder module initialized successfully.")
