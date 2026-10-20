"""
AgriSense: Create Dedicated Index, List of Figures, and List of Tables Document
Generates INDEX, LIST OF FIGURES, and LIST OF TABLES where page numbers (1, 2, 3, ...)
strictly start from Chapter 1: Introduction (Page 1).
Output: AgriSense_Index_and_List_of_Figures.docx
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "AgriSense_Index_and_List_of_Figures.docx")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=90, bottom=90, left=130, right=130):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0BEC5"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="8" w:space="0" w:color="2C3E50"/>
            <w:bottom w:val="single" w:sz="10" w:space="0" w:color="2C3E50"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_paragraph(p, space_before=0, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align

def add_heading_title(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=14, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1A, 0x25, 0x2C)
    return p

def create_index_document():
    doc = Document()

    # University standard margins: 1.25" Left (binding), 1.0" Right, Top, Bottom
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("AgriSense: Deep Learning Crop Disease Detection & Multilingual RAG Platform")
        hr.font.name = "Times New Roman"
        hr.font.size = Pt(8.5)
        hr.font.italic = True
        hr.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = fp.add_run("Department of CSE (AI & ML) - DKTE Society's TEI, Ichalkaranji")
        fr.font.name = "Times New Roman"
        fr.font.size = Pt(9)
        fr.font.color.rgb = RGBColor(0x95, 0xA5, 0xA6)

    # =========================================================================
    # 1. INDEX (TABLE OF CONTENTS)
    # Page numbers (1, 2, 3...) strictly start from Chapter 1: Introduction
    # =========================================================================
    add_heading_title(doc, "INDEX")

    index_headers = ["Sr. No.", "Contents", "Page No."]
    
    # Front matter in Roman numerals; Chapters start at Page 1
    index_data = [
        # Front Matter
        ["—", "Certificate", "i"],
        ["—", "Declaration", "ii"],
        ["—", "Acknowledgement", "iii"],
        ["—", "Abstract", "iv"],
        ["—", "List of Figures", "vi"],
        ["—", "List of Tables", "vii"],
        ["—", "List of Abbreviations & Mathematical Notations", "viii"],
        
        # Chapter 1 (Starts at Page 1)
        ["1.", "Introduction", "1"],
        ["", "1.1 Background & Motivation", "2"],
        ["", "1.2 Problem Statement", "3"],
        ["", "1.3 Objectives", "4"],
        ["", "1.4 Scope of the Project", "5"],
        
        # Chapter 2
        ["2.", "Literature Review", "7"],
        ["", "2.1 Classical & ML Baselines", "8"],
        ["", "2.2 Evolution of Deep Learning Architectures in the Domain", "8"],
        ["", "    2.2.1 Convolutional Neural Networks (CNNs) & Residual Learning", "8"],
        ["", "    2.2.2 Recurrent Neural Networks (RNNs) & Long Short-Term Memory (LSTM)", "9"],
        ["", "    2.2.3 Vision Transformers (ViTs) & Spatial Attention Mechanisms", "9"],
        ["", "    2.2.4 Generative & Diffusion Models in Agricultural Vision", "9"],
        ["", "    2.2.5 Multimodal Foundation Models & Retrieval-Augmented Generation", "10"],
        ["", "2.3 Literature Survey", "10"],
        ["", "2.4 Research Gap Identified", "13"],
        
        # Chapter 3
        ["3.", "Dataset Formulation & Preprocessing", "15"],
        ["", "3.1 Data Source & Acquisition", "16"],
        ["", "3.2 Data Preprocessing & Cleaning", "16"],
        ["", "3.3 Data Augmentation & Synthesis", "17"],
        ["", "3.4 Train-Validation-Test Splitting", "18"],
        
        # Chapter 4
        ["4.", "System Architecture & Proposed Methodology", "19"],
        ["", "4.1 Overall System Pipeline", "20"],
        ["", "4.2 Mathematical Formulation", "21"],
        ["", "    4.2.1 Deep Convolutional Classifier Formulation", "21"],
        ["", "    4.2.2 Agro-Meteorological Outbreak Risk Model", "22"],
        ["", "    4.2.3 Dense Vector Retrieval & Semantic Similarity", "23"],
        ["", "4.3 Training Strategy & Optimization", "24"],
        
        # Chapter 5
        ["5.", "Implementation Environment & Setup", "26"],
        ["", "5.1 Hardware Specifications", "27"],
        ["", "5.2 Software Stack & Frameworks", "27"],
        
        # Chapter 6
        ["6.", "Experimental Results, Evaluation & Analysis", "29"],
        ["", "6.1 Evaluation Metrics", "30"],
        ["", "6.2 Performance Comparison", "31"],
        ["", "6.3 Model Explainability & Interpretability", "33"],
        ["", "6.4 Error & Failure Analysis", "34"],
        
        # Chapter 7
        ["7.", "Conclusion & Future Scope", "36"],
        ["", "7.1 Summary of Contributions", "37"],
        ["", "7.2 Key Challenges Faced", "37"],
        ["", "7.3 Future Directions", "38"],
        
        # Chapter 8
        ["8.", "References", "39"]
    ]

    table_index = doc.add_table(rows=len(index_data) + 1, cols=3)
    table_index.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_index.autofit = False
    set_table_borders(table_index)

    col_widths = [0.9, 4.4, 0.9]

    # Format header
    hdr_cells = table_index.rows[0].cells
    for i, title in enumerate(index_headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "2C3E50")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Format data rows
    for r_idx, row in enumerate(index_data):
        row_cells = table_index.rows[r_idx + 1].cells
        is_chapter = bool(row[0] and row[0].endswith("."))
        bg_color = "EBF2FA" if is_chapter else ("F9FAFB" if r_idx % 2 == 1 else "FFFFFF")
        
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=70, bottom=70, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            format_paragraph(p, space_before=1, space_after=1, line_spacing=1.15)
            
            if c_idx == 0 or c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10 if not is_chapter else 10.5)
                if is_chapter:
                    r.font.bold = True

    for row in table_index.rows:
        for c_idx, w in enumerate(col_widths):
            row.cells[c_idx].width = Inches(w)

    doc.add_page_break()

    # =========================================================================
    # 2. LIST OF FIGURES
    # Page numbers strictly start from Chapter 1: Introduction (Page 1)
    # =========================================================================
    add_heading_title(doc, "LIST OF FIGURES")

    fig_headers = ["Figure No.", "Title of the Figure", "Page No."]
    
    # Corrected page numbers indexed from Introduction = Page 1
    fig_data = [
        ["Figure 1.1", "Functional block diagram contrasting traditional disconnected agricultural diagnostics with the unified AgriSense multimodal advisory paradigm", "3"],
        ["Figure 3.1", "Data preprocessing, stratified splitting, and augmentation pipeline for leaf disease classification", "17"],
        ["Figure 4.1", "Comprehensive end-to-end system architecture of the AgriSense platform", "21"],
        ["Figure 4.2", "Transfer learning convolutional architecture with custom dense classification head and dropout regularization", "22"],
        ["Figure 4.3", "Multi-source agro-meteorological disease risk engine and epidemiological rule processor", "23"],
        ["Figure 4.4", "Autonomous multi-tool multilingual Retrieval-Augmented Generation (RAG) architecture", "24"],
        ["Figure 6.1", "Training and validation loss and accuracy curves across 25 training epochs", "30"],
        ["Figure 6.2", "Normalized confusion matrix across 12 representative plant pathology classes", "32"],
        ["Figure 6.3", "Multi-architecture performance comparison (Accuracy, F1-Score, and CPU Latency)", "33"],
        ["Figure 6.4", "Latency profiling across system microservices and request processing pipelines", "35"]
    ]

    table_fig = doc.add_table(rows=len(fig_data) + 1, cols=3)
    table_fig.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_fig.autofit = False
    set_table_borders(table_fig)

    fig_widths = [1.3, 4.0, 0.9]

    # Format header
    hdr_cells = table_fig.rows[0].cells
    for i, title in enumerate(fig_headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "2C3E50")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Format data rows
    for r_idx, row in enumerate(fig_data):
        row_cells = table_fig.rows[r_idx + 1].cells
        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=75, bottom=75, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            format_paragraph(p, space_before=1, space_after=1, line_spacing=1.15)
            
            if c_idx == 0 or c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                if c_idx == 0:
                    r.font.bold = True

    for row in table_fig.rows:
        for c_idx, w in enumerate(fig_widths):
            row.cells[c_idx].width = Inches(w)

    doc.add_page_break()

    # =========================================================================
    # 3. LIST OF TABLES
    # Page numbers strictly start from Chapter 1: Introduction (Page 1)
    # =========================================================================
    add_heading_title(doc, "LIST OF TABLES")

    tbl_headers = ["Table No.", "Title of the Table", "Page No."]
    
    # Corrected page numbers indexed from Introduction = Page 1
    tbl_data = [
        ["Table 1.1", "Core System Objectives and Measurable Technical Performance Benchmarks", "5"],
        ["Table 2.1", "Comparative Literature Review of Modern Agricultural Diagnostic Systems", "12"],
        ["Table 2.2", "Identified Research Gaps and Technical Contributions of the AgriSense System", "14"],
        ["Table 3.1", "Stratified Distribution of the Plant Pathology Benchmark Dataset Across Crops and Partitions", "18"],
        ["Table 5.1", "Hardware Environment Specifications for Training and Local Deployment", "27"],
        ["Table 5.2", "Software Stack, Frameworks, and Dependent Runtime Libraries", "28"],
        ["Table 6.1", "Quantitative Performance Benchmark Across Evaluated Deep Learning Model Architectures", "31"],
        ["Table 7.1", "Validation of Measurable Outcomes Against Chapter 1 Defined Project Objectives", "37"]
    ]

    table_tbl = doc.add_table(rows=len(tbl_data) + 1, cols=3)
    table_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_tbl.autofit = False
    set_table_borders(table_tbl)

    tbl_widths = [1.3, 4.0, 0.9]

    # Format header
    hdr_cells = table_tbl.rows[0].cells
    for i, title in enumerate(tbl_headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "2C3E50")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Format data rows
    for r_idx, row in enumerate(tbl_data):
        row_cells = table_tbl.rows[r_idx + 1].cells
        bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=75, bottom=75, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            format_paragraph(p, space_before=1, space_after=1, line_spacing=1.15)
            
            if c_idx == 0 or c_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                if c_idx == 0:
                    r.font.bold = True

    for row in table_tbl.rows:
        for c_idx, w in enumerate(tbl_widths):
            row.cells[c_idx].width = Inches(w)

    print(f"[*] Saving updated Index and List of Figures document to {OUTPUT_PATH}...")
    try:
        doc.save(OUTPUT_PATH)
        print(f"[OK] Successfully built {OUTPUT_PATH} (Size: {os.path.getsize(OUTPUT_PATH)} bytes)")
    except PermissionError:
        alt_path = os.path.join(os.path.dirname(__file__), "AgriSense_Index_and_List_of_Figures_Updated.docx")
        print(f"[!] {OUTPUT_PATH} is currently open in Word. Saving to alternate file: {alt_path}")
        doc.save(alt_path)
        print(f"[OK] Successfully built {alt_path} (Size: {os.path.getsize(alt_path)} bytes)")

if __name__ == "__main__":
    create_index_document()
